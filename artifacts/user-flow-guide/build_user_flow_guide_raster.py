from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_user_flow_guide import ERROR_ROWS, SCREENS


ROOT = Path(__file__).resolve().parent
SHOT_DIR = ROOT / "screenshots"
PAGE_DIR = ROOT / "composed_pages"
OUTPUT = ROOT / "StoryWalk_用户使用动线与错误提示演示指南.docx"

PAGE_W, PAGE_H = 2550, 3300
LEFT, RIGHT = 230, 2320
CONTENT_W = RIGHT - LEFT
FONT_LIGHT = "/System/Library/Fonts/STHeiti Light.ttc"
FONT_MEDIUM = "/System/Library/Fonts/STHeiti Medium.ttc"

COLORS = {
    "ink": "#20252B",
    "muted": "#5F6B76",
    "blue": "#2E74B5",
    "dark_blue": "#1F4D78",
    "pale_blue": "#E8EEF5",
    "pale_gray": "#F2F4F7",
    "pale_red": "#FBEAEC",
    "red": "#9B1C1C",
    "line": "#CCD5DF",
    "white": "#FFFFFF",
}


def font(size: int, bold: bool = False):
    return ImageFont.truetype(FONT_MEDIUM if bold else FONT_LIGHT, size=size, index=0)


def wrap_text(draw: ImageDraw.ImageDraw, text: str, selected_font, max_width: int) -> list[str]:
    lines: list[str] = []
    current = ""
    for char in text:
        candidate = current + char
        if char == "\n":
            lines.append(current)
            current = ""
            continue
        if draw.textbbox((0, 0), candidate, font=selected_font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current.rstrip())
            current = char.lstrip()
    if current or not lines:
        lines.append(current)
    return lines


def draw_wrapped(draw, xy, text, selected_font, fill, max_width, line_gap=12, max_lines=None):
    x, y = xy
    lines = wrap_text(draw, text, selected_font, max_width)
    if max_lines is not None:
        lines = lines[:max_lines]
    bbox = draw.textbbox((0, 0), "国Ag", font=selected_font)
    line_height = bbox[3] - bbox[1] + line_gap
    for line in lines:
        draw.text((x, y), line, font=selected_font, fill=fill)
        y += line_height
    return y


def base_page(page_number: int | None = None):
    image = Image.new("RGB", (PAGE_W, PAGE_H), COLORS["white"])
    draw = ImageDraw.Draw(image)
    if page_number is not None:
        draw.text((LEFT, 95), "STORYWALK · 产品操作指南", font=font(27, True), fill=COLORS["muted"])
        footer = f"StoryWalk  |  {page_number}"
        box = draw.textbbox((0, 0), footer, font=font(25))
        draw.text((RIGHT - (box[2] - box[0]), 3205), footer, font=font(25), fill=COLORS["muted"])
    return image, draw


def paste_contained(canvas, source_path: Path, box: tuple[int, int, int, int]):
    x1, y1, x2, y2 = box
    source = Image.open(source_path).convert("RGB")
    source.thumbnail((x2 - x1, y2 - y1), Image.Resampling.LANCZOS)
    x = x1 + (x2 - x1 - source.width) // 2
    y = y1 + (y2 - y1 - source.height) // 2
    canvas.paste(source, (x, y))
    return (x, y, x + source.width, y + source.height)


def save_page(image: Image.Image, page_number: int) -> Path:
    PAGE_DIR.mkdir(parents=True, exist_ok=True)
    path = PAGE_DIR / f"page-{page_number:02d}.png"
    image.save(path, "PNG", optimize=True)
    return path


def cover_page() -> Image.Image:
    image, draw = base_page()
    draw.text((LEFT, 250), "产品演示手册", font=font(31, True), fill=COLORS["blue"])
    draw.text((LEFT, 365), "StoryWalk 用户使用动线", font=font(73, True), fill=COLORS["ink"])
    draw.text((LEFT, 495), "逐页截图、操作说明与错误提示触发指南", font=font(39), fill=COLORS["dark_blue"])
    paste_contained(image, SHOT_DIR / "01-language-home.png", (LEFT, 700, RIGHT, 1820))
    draw.text((LEFT, 1920), "用途：", font=font(30, True), fill=COLORS["ink"])
    draw.text((LEFT + 115, 1920), "项目演示、用户验收、答辩讲解与测试复现", font=font(30), fill=COLORS["ink"])
    draw.text((LEFT, 1990), "截图环境：", font=font(30, True), fill=COLORS["ink"])
    draw.text((LEFT + 160, 1990), "Google Chrome · 本地开发环境 · 2026-08-31", font=font(30), fill=COLORS["ink"])
    draw.rounded_rectangle((LEFT, 2145, RIGHT, 2325), radius=18, fill=COLORS["pale_blue"])
    draw.text((LEFT + 32, 2182), "隐私说明", font=font(29, True), fill=COLORS["dark_blue"])
    draw.text((LEFT + 200, 2182), "文档不包含真实密码或安全问题答案；错误截图使用虚构测试数据。", font=font(29), fill=COLORS["ink"])
    return image


def overview_page(page_number: int) -> Image.Image:
    image, draw = base_page(page_number)
    draw.text((LEFT, 210), "用户主线总览", font=font(52, True), fill=COLORS["blue"])
    intro = "本指南按一次完整旅行体验组织。演示时可沿下列顺序讲解；每个页面后的“下一步”可直接作为口播衔接。"
    y = draw_wrapped(draw, (LEFT, 320), intro, font(30), COLORS["ink"], CONTENT_W, 15)
    flow = [
        "选择语言，决定以游客方式开始或登录账号。",
        "登录/注册，并在注册时设置安全问题。",
        "通过 AI 对话或手动选项确认旅行偏好与到达日期。",
        "生成路线，按站点推进并使用景点导览或 StoryWalk。",
        "在个人中心维护密码、安全问题与长期偏好。",
        "用回忆录整理真实照片、章节和分享隐私。",
        "用明信片保存每个到访地点，并处理下载、分享或重新生成。",
    ]
    y += 45
    for idx, item in enumerate(flow, 1):
        draw.ellipse((LEFT, y + 2, LEFT + 54, y + 56), fill=COLORS["blue"])
        number = str(idx)
        box = draw.textbbox((0, 0), number, font=font(26, True))
        draw.text((LEFT + 27 - (box[2] - box[0]) / 2, y + 7), number, font=font(26, True), fill=COLORS["white"])
        y = draw_wrapped(draw, (LEFT + 85, y), item, font(30), COLORS["ink"], CONTENT_W - 85, 13) + 24
    y += 20
    draw.text((LEFT, y), "截图阅读方式", font=font(39, True), fill=COLORS["blue"])
    y += 78
    labels = [
        ("页面目的", "为什么用户会来到此页。"),
        ("用户动作", "演示时应执行或说明的核心操作。"),
        ("下一步", "完成当前动作后的自然去向。"),
        ("错误触发", "出现提示的条件及建议恢复方式；无需实际修改演示账号。"),
    ]
    for label, value in labels:
        draw.text((LEFT, y), label, font=font(29, True), fill=COLORS["dark_blue"])
        draw.text((LEFT + 175, y), value, font=font(29), fill=COLORS["ink"])
        y += 66
    return image


def screenshot_page(item: dict[str, str], step: int, page_number: int) -> Image.Image:
    image, draw = base_page(page_number)
    draw.text((LEFT, 185), f"{item['group']}  ·  STEP {step:02d}", font=font(28, True), fill=COLORS["blue"])
    draw.text((LEFT, 255), item["title"], font=font(50, True), fill=COLORS["blue"])
    placed = paste_contained(image, SHOT_DIR / item["file"], (LEFT, 390, RIGHT, 1600))
    caption = f"图 {step}  {item['title']}（Chrome 实际页面）"
    cap_box = draw.textbbox((0, 0), caption, font=font(24))
    draw.text(((PAGE_W - (cap_box[2] - cap_box[0])) / 2, placed[3] + 22), caption, font=font(24), fill=COLORS["muted"])

    y = max(1710, placed[3] + 90)
    fields = (("页面目的", item["purpose"]), ("用户动作", item["action"]), ("下一步", item["next"]))
    for label, value in fields:
        draw.text((LEFT, y), label, font=font(29, True), fill=COLORS["dark_blue"])
        y = draw_wrapped(draw, (LEFT + 190, y), value, font(29), COLORS["ink"], CONTENT_W - 190, 12) + 26

    if item.get("error"):
        lines = wrap_text(draw, item["error"], font(28), CONTENT_W - 250)
        box_h = 48 + len(lines) * 45
        draw.rounded_rectangle((LEFT, y + 5, RIGHT, y + box_h), radius=16, fill=COLORS["pale_red"])
        draw.text((LEFT + 28, y + 28), "错误触发", font=font(28, True), fill=COLORS["red"])
        line_y = y + 28
        for line in lines:
            draw.text((LEFT + 210, line_y), line, font=font(28), fill=COLORS["ink"])
            line_y += 45
    return image


def error_table_pages(start_page_number: int) -> list[Image.Image]:
    pages: list[Image.Image] = []
    chunks = [ERROR_ROWS[:8], ERROR_ROWS[8:]]
    for part, rows in enumerate(chunks, 1):
        page_number = start_page_number + part - 1
        image, draw = base_page(page_number)
        title = "错误提示触发表" if part == 1 else "错误提示触发表（续）"
        draw.text((LEFT, 205), title, font=font(50, True), fill=COLORS["blue"])
        if part == 1:
            intro = "优先在虚构测试资料或空表单上复现；不要为了截图反复尝试真实账号密码或安全答案。"
            draw_wrapped(draw, (LEFT, 300), intro, font(28), COLORS["ink"], CONTENT_W, 12)
            y = 425
        else:
            y = 335

        widths = [250, 500, 760, 580]
        x_positions = [LEFT]
        for width in widths:
            x_positions.append(x_positions[-1] + width)
        header_h = 92
        draw.rectangle((LEFT, y, RIGHT, y + header_h), fill=COLORS["pale_blue"], outline=COLORS["line"], width=3)
        headers = ["区域", "触发条件", "用户看到的提示", "建议恢复"]
        for col, label in enumerate(headers):
            draw.text((x_positions[col] + 16, y + 26), label, font=font(25, True), fill=COLORS["dark_blue"])
            if col:
                draw.line((x_positions[col], y, x_positions[col], y + header_h), fill=COLORS["line"], width=3)
        y += header_h

        for row_idx, row in enumerate(rows):
            row_fonts = [font(23, True), font(23), font(23), font(23)]
            wrapped = [wrap_text(draw, value, row_fonts[i], widths[i] - 30) for i, value in enumerate(row)]
            row_h = max(105, max(len(lines) for lines in wrapped) * 39 + 34)
            fill = COLORS["pale_gray"] if row_idx % 2 else COLORS["white"]
            draw.rectangle((LEFT, y, RIGHT, y + row_h), fill=fill, outline=COLORS["line"], width=3)
            for col, lines in enumerate(wrapped):
                if col:
                    draw.line((x_positions[col], y, x_positions[col], y + row_h), fill=COLORS["line"], width=3)
                line_y = y + 18
                for line in lines:
                    draw.text((x_positions[col] + 15, line_y), line, font=row_fonts[col], fill=COLORS["ink"])
                    line_y += 39
            y += row_h

        if part == 2:
            y += 70
            draw.text((LEFT, y), "演示注意事项", font=font(39, True), fill=COLORS["blue"])
            y += 85
            notes = [
                "注册重复邮箱会明确提示更换邮箱或直接登录；不要创建重复数据。",
                "修改密码的最终提交会改变账号密码，演示前应确认。",
                "安全问题答案不会明文展示；找回密码时仅显示问题文本。",
                "天气、地图或图片失败时，应保留用户选择并提供重试或降级说明。",
                "删除明信片、撤销分享链接等操作不建议在主演示账号上执行。",
            ]
            for note in notes:
                draw.ellipse((LEFT, y + 12, LEFT + 14, y + 26), fill=COLORS["blue"])
                y = draw_wrapped(draw, (LEFT + 36, y), note, font(27), COLORS["ink"], CONTENT_W - 36, 11) + 18
        pages.append(image)
    return pages


def make_docx(page_paths: list[Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.025)
    section.right_margin = Inches(0.025)
    section.top_margin = Inches(0.025)
    section.bottom_margin = Inches(0.025)
    section.header_distance = Inches(0)
    section.footer_distance = Inches(0)
    style = doc.styles["Normal"]
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing = 1

    for idx, page_path in enumerate(page_paths):
        paragraph = doc.add_paragraph()
        if idx:
            paragraph.paragraph_format.page_break_before = True
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run()
        inline = run.add_picture(str(page_path), width=Inches(8.43))
        inline._inline.docPr.set("descr", f"StoryWalk 用户动线指南第 {idx + 1} 页")

    props = doc.core_properties
    props.title = "StoryWalk 用户使用动线与错误提示演示指南"
    props.subject = "StoryWalk 产品逐页操作截图与错误触发说明"
    props.author = "StoryWalk Project Team"
    doc.save(OUTPUT)


def build() -> Path:
    page_paths: list[Path] = []
    page_paths.append(save_page(cover_page(), 1))
    page_paths.append(save_page(overview_page(2), 2))
    for step, item in enumerate(SCREENS, 1):
        page_number = len(page_paths) + 1
        page_paths.append(save_page(screenshot_page(item, step, page_number), page_number))
    for image in error_table_pages(len(page_paths) + 1):
        page_number = len(page_paths) + 1
        page_paths.append(save_page(image, page_number))
    make_docx(page_paths)
    print(f"pages={len(page_paths)} output={OUTPUT}")
    return OUTPUT


if __name__ == "__main__":
    build()
